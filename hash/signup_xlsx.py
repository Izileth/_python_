import hashlib
import os
from openpyxl import Workbook, load_workbook
from pyfiglet import figlet_format
from rich.console import Console
from rich.panel import Panel

console = Console()

# ==========================================
#                REGISTER
# ==========================================

console.print(
    figlet_format("REGISTER", font="slant"),
    style="bold cyan"
)

console.print("[dim]Create a new account[/dim]\n")

# ==========================================
#                DATABASE
# ==========================================

ARCHIVE = "database.xlsx"

def load_users():
    """Carrega os usuários do Excel. Retorna um dicionário."""
    if not os.path.exists(ARCHIVE):
        return {}

    wb = load_workbook(ARCHIVE)
    ws = wb.active

    users = {}
    # Pula o cabeçalho (linha 1)
    for row in ws.iter_rows(min_row=2, values_only=True):
        username, salt, password_hash = row
        if username:
            users[username] = {
                "salt": salt,
                "hash": password_hash
            }
    return users

def save_users(users):
    """Salva o dicionário de usuários no Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    # Cabeçalho
    ws.append(["username", "salt", "hash"])

    # Dados
    for username, data in users.items():
        ws.append([username, data["salt"], data["hash"]])

    # Ajusta largura das colunas
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 70

    wb.save(ARCHIVE)

# ==========================================
#            PASSWORD HASHING
# ==========================================

def hash_password(password):
    salt = os.urandom(16)

    password_hash = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        100_000
    )

    return {
        "salt": salt.hex(),
        "hash": password_hash.hex()
    }

# ==========================================
#              CREATE USER
# ==========================================

users = load_users()

user = console.input("[cyan]Create a user:[/cyan] ").strip()
password = console.input("[cyan]Create a password:[/cyan] ")
confirm_password = console.input("[cyan]Confirm password:[/cyan] ")

# ==========================================
#          PASSWORD VALIDATION
# ==========================================

if password != confirm_password:
    console.print(
        Panel(
            "[bold red]✗ Passwords do not match.[/bold red]",
            title="REGISTER",
            border_style="red"
        )
    )
    exit()

# ==========================================
#          USER VALIDATION
# ==========================================

if user in users:
    console.print(
        Panel(
            "[bold red]✗ User already exists.[/bold red]",
            title="REGISTER",
            border_style="red"
        )
    )
    exit()

if not user:
    console.print(
        Panel(
            "[bold red]✗ Username cannot be empty.[/bold red]",
            title="REGISTER",
            border_style="red"
        )
    )
    exit()

# ==========================================
#              SAVE USER
# ==========================================

users[user] = hash_password(password)
save_users(users)

# ==========================================
#               SUCCESS
# ==========================================

console.print(
    Panel(
        f"[bold green]✓ User '{user}' created successfully![/bold green]\n[dim]Saved to database.xlsx[/dim]",
        title="REGISTER",
        border_style="green"
    )
)