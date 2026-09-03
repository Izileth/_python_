import hashlib
import os
from openpyxl import load_workbook
from pyfiglet import figlet_format
from rich.console import Console
from rich.panel import Panel

console = Console()

# ==========================================
#                  LOGIN
# ==========================================

console.print(
    figlet_format("LOGIN", font="slant"),
    style="bold cyan"
)

console.print("[dim]Enter your credentials[/dim]\n")

# ==========================================
#                DATABASE
# ==========================================

ARCHIVE = "database.xlsx"

def load_users():
    """Carrega os usuários do Excel. Retorna um dicionário."""
    if not os.path.exists(ARCHIVE):
        return {}

    wb = load_workbook(ARCHIVE)
    ws = wb.active

    users = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        username, salt, password_hash = row
        if username:
            users[username] = {
                "salt": salt,
                "hash": password_hash
            }
    return users

# ==========================================
#          PASSWORD VERIFICATION
# ==========================================

def verify_password(password: str, salt_hex: str, stored_hash: str) -> bool:
    """Verifica se a senha está correta usando o salt salvo."""
    salt = bytes.fromhex(salt_hex)

    password_hash = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        100_000
    )

    return password_hash.hex() == stored_hash

# ==========================================
#              LOGIN PROCESS
# ==========================================

users = load_users()

if not users:
    console.print(
        Panel(
            "[bold red]✗ No users found in the database.[/bold red]\n[dim]Please register first.[/dim]",
            title="LOGIN",
            border_style="red"
        )
    )
    exit()

username = console.input("[cyan]Username:[/cyan] ").strip()
password = console.input("[cyan]Password:[/cyan] ")

# ==========================================
#              VALIDATIONS
# ==========================================

if username not in users:
    console.print(
        Panel(
            "[bold red]✗ User not found.[/bold red]",
            title="LOGIN",
            border_style="red"
        )
    )
    exit()

user_data = users[username]

if not verify_password(password, user_data["salt"], user_data["hash"]):
    console.print(
        Panel(
            "[bold red]✗ Incorrect password.[/bold red]",
            title="LOGIN",
            border_style="red"
        )
    )
    exit()

# ==========================================
#               SUCCESS
# ==========================================

console.print(
    Panel(
        f"[bold green]✓ Welcome back, {username}![/bold green]\n[dim]Login successful.[/dim]",
        title="LOGIN",
        border_style="green"
    )
)